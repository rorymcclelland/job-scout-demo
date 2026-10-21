"""Writes the demo candidate's base CV and covering letter (H16). Run: python demo/make_documents.py

The pack builder (tools/build_pack.py) edits copies of these exactly as it edits the
owner's own documents: the Summary, the three highlights and the location line on the
CV, and the date, salutation and body of the letter. All made up."""
import os
from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt, Cm

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "documents")
NAME, CONTACT = "Sam Calder", "sam.calder@example.com  ·  +44 7700 900000"


def base():
    d = Document()
    for s in d.sections:
        s.top_margin = s.bottom_margin = Cm(1.8)
        s.left_margin = s.right_margin = Cm(2.0)
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    return d


def heading(d, text):
    # A real heading style, as in the owner's own CV: the pack builder's Markdown copy, and so the
    # reviewer's sentence list, find the Summary and highlights by it (found by the first demo run).
    p = d.add_paragraph(style="Heading 1")
    r = p.add_run(text.upper())
    r.bold = True
    r.font.size = Pt(11)
    return p


def role(d, title, dates, bullets, page_break=False):
    p = d.add_paragraph()
    if page_break:
        p.add_run().add_break(WD_BREAK.PAGE)
    r = p.add_run(title)
    r.bold = True
    p.add_run(f"  ·  {dates}")
    for b in bullets:
        d.add_paragraph(b, style="List Bullet")


def cv():
    d = base()
    p = d.add_paragraph(); r = p.add_run(NAME); r.bold = True; r.font.size = Pt(20)
    d.add_paragraph(f"London  ·  {CONTACT}")
    heading(d, "Summary")
    d.add_paragraph("Product leader with eighteen years in consumer and B2B software, building products that put people in control of their own data.")
    heading(d, "Selected highlights")
    for h in ["Redesigned a browser's permissions prompt used by hundreds of millions of people.",
              "Took a digital asset management platform to market as multi-tenant SaaS on a private cloud.",
              "Led discovery, bids and pilots using AI for consultancy clients."]:
        d.add_paragraph(h, style="List Bullet")
    heading(d, "Experience")
    role(d, "Product Director, Fieldwork Studio", "2024 to date", [
        "Led discovery for clients exploring AI: what to build, what to buy and what not to do.",
        "Wrote and led bids; ran two pilots of AI assistants with client teams.",
        "Builds working prototypes with AI coding tools to test ideas before a build is funded."])
    role(d, "Head of Product, Lingua", "2021 to 2023", [
        "Owned the learning experience of a language-learning app: onboarding, the daily lesson and the paid tier.",
        "Rebuilt onboarding around a short placement test; first-week completion rose from 31% to 38%."])
    role(d, "Interim product plan, Meridian Members", "2023", [
        "A three-month product plan for a members' club group, with projected membership and revenue figures."])
    role(d, "Group Product Manager, Privacy and Permissions, Northwind Browser", "2016 to 2021", [
        "Owned the privacy settings and the permissions prompt for a browser used by hundreds of millions of people.",
        "Redesigned the permissions prompt so a site asks only when the person acts; prompts shown fell by about a third.",
        "Ran each launch through privacy review with legal and two data-protection regulators.",
        "Led four product managers working with about forty engineers across two sites."], page_break=True)
    role(d, "Product Lead, then Head of Product, Halcyon Assets", "2008 to 2016", [
        "Took the platform to market as multi-tenant SaaS on a private cloud, moving customers over eighteen months.",
        "Built integrations with three enterprise systems; grew the product team from two to nine."])
    heading(d, "Earlier")
    d.add_paragraph("Taught English in South Korea on a government exchange programme for two years. Degree in mechanical engineering.")
    d.save(os.path.join(OUT, f"{NAME} - CV.docx"))


def letter():
    d = base()
    p = d.add_paragraph(); r = p.add_run(NAME); r.bold = True; r.font.size = Pt(16)
    d.add_paragraph(f"London  ·  {CONTACT}")
    d.add_paragraph("20 October 2026")
    d.add_paragraph("Dear Hiring Manager,")
    for t in ["I am writing to apply for the role.", "Second paragraph.", "Third paragraph."]:
        d.add_paragraph(t)
    d.add_paragraph("Kind regards,")
    d.add_paragraph(NAME)
    d.save(os.path.join(OUT, f"{NAME} - covering letter.docx"))


os.makedirs(OUT, exist_ok=True)
cv()
letter()
print("ok", sorted(os.listdir(OUT)))
