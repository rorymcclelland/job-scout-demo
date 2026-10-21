"""Tests for the letter word-count check in build_pack.py (backlog E14)."""

import unittest

from build_pack import check_letter_words


def letter(n):
    """A letter body of exactly n words, over three paragraphs."""
    words = ["word"] * n
    third = n // 3
    return [" ".join(words[:third]), " ".join(words[third:2 * third]), " ".join(words[2 * third:])]


RANGE = {"min": 350, "max": 450}


class LetterWords(unittest.TestCase):
    def test_one_word_short_fails_and_says_so(self):
        n, problem, sign_off = check_letter_words(letter(349), RANGE)
        self.assertEqual(n, 349)
        self.assertIn("1 word short of 350", problem)
        self.assertIsNone(sign_off)

    def test_both_ends_of_the_range_pass(self):
        for n in (350, 450):
            count, problem, sign_off = check_letter_words(letter(n), RANGE)
            self.assertEqual(count, n)
            self.assertIsNone(problem)
            self.assertIsNone(sign_off)

    def test_one_word_over_fails(self):
        _, problem, _ = check_letter_words(letter(451), RANGE)
        self.assertIn("1 word over 450", problem)

    def test_missing_range_falls_back_to_350_to_450(self):
        self.assertIsNotNone(check_letter_words(letter(320))[1])
        self.assertIsNone(check_letter_words(letter(400))[1])
        self.assertIsNone(check_letter_words(letter(400), {"min": "x"})[1])

    def test_range_from_settings_is_used(self):
        self.assertIsNone(check_letter_words(letter(320), {"min": 300, "max": 400})[1])

    def test_after_one_revision_it_goes_to_sign_off_not_a_failure(self):
        n, problem, sign_off = check_letter_words(letter(330), RANGE, revised=True)
        self.assertIsNone(problem)
        self.assertIn("20 words short of 350", sign_off)
        self.assertIn("after one revision", sign_off)


if __name__ == "__main__":
    unittest.main()


class SpecValues(unittest.TestCase):
    """Sprint 8 review: the demo's name, contact and page-two values must not be empty or a path."""

    def run_spec(self, **over):
        import json, os, subprocess, sys, tempfile
        with tempfile.TemporaryDirectory() as d:
            spec = {"outDir": d, "label": "x", **over}
            p = os.path.join(d, "spec.json")
            with open(p, "w", encoding="utf-8") as f:
                json.dump(spec, f)
            return subprocess.run([sys.executable, os.path.join(os.path.dirname(__file__), "build_pack.py"), p], capture_output=True, text=True)

    def test_empty_contact_refused(self):
        r = self.run_spec(name="Sam Calder", contact="")
        self.assertNotEqual(r.returncode, 0)
        self.assertIn("contact is empty", r.stderr)

    def test_name_with_a_path_refused(self):
        r = self.run_spec(name="..\\..\\Sam", contact="x", pageTwoStartsWith="Group Product Manager")
        self.assertNotEqual(r.returncode, 0)
        self.assertIn("file-name character", r.stderr)


class MarkdownCopy(unittest.TestCase):
    """Found by the first demo run: Word's List Bullet style must become a Markdown bullet."""

    def test_list_bullet_and_heading_become_markdown(self):
        import os, tempfile
        from docx import Document
        from build_pack import docx_to_md
        with tempfile.TemporaryDirectory() as d:
            doc = Document()
            doc.add_paragraph("SUMMARY", style="Heading 1")
            doc.add_paragraph("Product leader.")
            doc.add_paragraph("Did a thing.", style="List Bullet")
            p = os.path.join(d, "x.docx")
            doc.save(p)
            md = docx_to_md(p)
        self.assertIn("## Summary", md)
        self.assertIn("- Did a thing.", md)


class NoPersonalDefaults(unittest.TestCase):
    """H22: with no name in the spec, the build fails rather than using anyone's details."""

    def test_missing_name_refused(self):
        r = SpecValues.run_spec(SpecValues(), contact="x", pageTwoStartsWith="y")
        self.assertNotEqual(r.returncode, 0)
        self.assertIn("name is empty or missing", r.stderr)
