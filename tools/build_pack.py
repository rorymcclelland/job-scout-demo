"""Builds the Word and PDF files of an application pack (backlog E3, E4, E5).

Run by Claude during a Job Scout run, in its own workspace, after it has
written the tailored words into a spec file. The pack's look always comes from
the candidate's own documents, never from this script:

  CV      a copy of the candidate's clean two-page CV ("cvBase"). Only the
          Summary and the three Selected Highlights are replaced, plus the
          location on the contact line.
  Letter  a copy of a previous covering letter ("letterBase"), whose header and
          body styling are kept. Its date, salutation and body paragraphs are
          replaced.

Checks, which fail the build rather than produce a bad pack (G1):
  - the CV is exactly two pages, and page two starts with the spec's
    "pageTwoStartsWith" (the first entry on page two);
  - the letter is one page;
  - the letter's body is within the word range in the board's Settings
    (E14; 350 to 450 if the spec gives none). On a rebuild after one revision
    ("letterRevised": true) a count still out of range no longer fails the
    build: it is listed under "signOff" for the owner instead;
  - no em dashes, and none of the barred phrases (BARRED below).

Usage:
  python3 build_pack.py spec.json

spec.json:
  {
    "cvBase": "...Sam Calder - CV.docx",
    "letterBase": "...Sam Calder - covering letter.docx",
    "outDir": "...",                       # where the files are written
    "label": "Resonant Security",          # used in file names
    "name": "Sam Calder",                  # the candidate, as in the base documents (required)
    "contact": "sam.calder@example.com  ·  +44 7700 900000",  # the contact line after the place (required)
    "pageTwoStartsWith": "Group Product Manager, ...",          # page two's first entry (required)
    "location": "London",                  # or another city, or "" for none
    "summary": "...",
    "highlights": ["...", "...", "..."],
    "letter": {"date": "15 October 2026", "salutation": "Dear ...,", "paragraphs": ["...", "..."]},
    "letterWords": {"min": 350, "max": 450}, # from Settings, packs.letterMinWords and letterMaxWords
    "letterRevised": false                  # true on the rebuild after one revision
  }
"""

import copy
import json
import re
import subprocess
import sys
from pathlib import Path

try:
    from docx import Document
except ImportError:  # the word-count check is tested without python-docx installed
    Document = None

BARRED = [
    "RAG", "semantic vector search", "local LLM orchestration", "LLM-driven intent classification",
    "production AI systems built personally", "genuinely", "genuine", "quietly", "robust",
]


def set_text(paragraph, text):
    """Replace a paragraph's words, keeping the first run's formatting."""
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def set_location(doc, location, contact):
    for p in doc.paragraphs[:3]:
        if contact in p.text:
            rest = p.text[p.text.index(contact):]
            set_text(p, f"{location}  ·  {rest}" if location else rest)
            return
    raise SystemExit("Contact line not found in the base document.")


def index_after(doc, heading):
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().upper() == heading:
            return i + 1
    raise SystemExit(f"Heading {heading!r} not found in the CV base.")


def build_cv(spec, out_docx):
    doc = Document(spec["cvBase"])
    set_location(doc, spec.get("location", ""), spec["contact"])
    i = index_after(doc, "SUMMARY")
    set_text(doc.paragraphs[i], spec["summary"])
    j = index_after(doc, "SELECTED HIGHLIGHTS")
    highlights = spec["highlights"]
    if len(highlights) != 3:
        raise SystemExit("Exactly three highlights: page one overflows at four.")
    for k, h in enumerate(highlights):
        set_text(doc.paragraphs[j + k], h)
    if doc.paragraphs[j + 3].style is None or doc.paragraphs[j + 3].text.strip().upper() != "EXPERIENCE":
        raise SystemExit("The base CV does not have exactly three highlights; check the master.")
    doc.save(out_docx)


def build_letter(spec, out_docx):
    doc = Document(spec["letterBase"])
    set_location(doc, spec.get("location", ""), spec["contact"])
    paras = doc.paragraphs
    # Base layout: 0 name, 1 contact, 2 date, 3 salutation, 4..n-3 body, n-2 sign-off, n-1 name.
    letter = spec["letter"]
    set_text(paras[2], letter["date"])
    set_text(paras[3], letter["salutation"])
    body = paras[4:-2]
    template = copy.deepcopy(body[0]._p)
    for p in body:
        p._p.getparent().remove(p._p)
    insert_before = paras[-2]._p
    for text in letter["paragraphs"]:
        el = copy.deepcopy(template)
        insert_before.addprevious(el)
    doc.save(out_docx)
    # Re-open to fill the cloned paragraphs through python-docx.
    doc = Document(out_docx)
    for k, text in enumerate(letter["paragraphs"]):
        set_text(doc.paragraphs[4 + k], text)
    doc.save(out_docx)


def to_pdf(docx_path):
    out = docx_path.parent
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(out), str(docx_path)],
                   check=True, capture_output=True, timeout=180)
    return docx_path.with_suffix(".pdf")


def pdf_pages(pdf):
    text = subprocess.run(["pdftotext", "-layout", str(pdf), "-"], check=True, capture_output=True, text=True).stdout
    return text.split("\f")[:-1] if text.endswith("\f") else text.split("\f")


def check_words(label, texts):
    problems = []
    joined = "\n".join(texts)
    if "—" in joined:
        problems.append(f"{label}: contains an em dash.")
    for phrase in BARRED:
        if re.search(rf"\b{re.escape(phrase)}\b", joined, flags=re.IGNORECASE):
            problems.append(f"{label}: contains the barred word or claim '{phrase}'.")
    return problems


DEFAULT_WORDS = (350, 450)


def word_count(paragraphs):
    return sum(len(p.split()) for p in paragraphs)


def check_letter_words(paragraphs, words=None, revised=False):
    """E14. Returns (count, problem or None, sign-off note or None)."""
    words = words or {}
    low = words.get("min", DEFAULT_WORDS[0]) if isinstance(words.get("min"), int) else DEFAULT_WORDS[0]
    high = words.get("max", DEFAULT_WORDS[1]) if isinstance(words.get("max"), int) else DEFAULT_WORDS[1]
    n = word_count(paragraphs)
    if low <= n <= high:
        return n, None, None
    gap = low - n if n < low else n - high
    how = f"{gap} word{'' if gap == 1 else 's'} {'short of' if n < low else 'over'} {low if n < low else high}"
    message = f"Letter is {n} words, {how} (range {low} to {high})"
    if revised:
        return n, None, f"{message}, after one revision: check the length before sending."
    return n, f"{message}: revise it once and rebuild with \"letterRevised\": true.", None


def paragraphs(docx_path):
    return [p.text.strip() for p in Document(docx_path).paragraphs if p.text.strip()]


def docx_to_md(docx_path):
    doc = Document(docx_path)
    lines = []
    for p in doc.paragraphs:
        name = p.style.name if p.style is not None else ""
        t = p.text.strip()
        if not t:
            continue
        if name == "Heading 1":
            lines += ["", f"## {t.title()}", ""]
        elif name == "Heading 2":
            lines += ["", f"### {t}"]
        elif name in ("List Paragraph", "List Bullet"):
            lines.append(f"- {t}")
        else:
            lines += [t, ""] if len(t) > 60 else [t]
    return "\n".join(lines).strip() + "\n"


def main():
    if len(sys.argv) != 2 or sys.argv[1] in ("-h", "--help"):
        print(__doc__)
        sys.exit(0 if len(sys.argv) == 2 else 2)
    spec = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    out = Path(spec["outDir"])
    out.mkdir(parents=True, exist_ok=True)
    label = spec["label"]
    # H16 (Sprint 8): the demo candidate's documents name their own owner, contact line and page-two entry.
    # H22: no personal defaults. The spec names the candidate, their contact line and page two's first entry.
    name = spec.get("name")
    page_two = spec.get("pageTwoStartsWith")
    # Review, Sprint 8: an empty value would match the wrong paragraph, and a name becomes part of a file name.
    for key, value in (("name", name), ("contact", spec.get("contact")), ("pageTwoStartsWith", page_two)):
        if not isinstance(value, str) or not value.strip():
            sys.exit(f"build_pack: the spec's {key} is empty or missing (H22: there are no defaults).")
    if re.search(r'[\\/:*?"<>|]', name) or ".." in name:
        sys.exit("build_pack: the spec's name cannot contain a path or file-name character.")
    cv = out / f"{name} - CV - {label}.docx"
    letter = out / f"{name} - covering letter - {label}.docx"

    problems = check_words("CV", [spec["summary"], *spec["highlights"]])
    problems += check_words("Letter", spec["letter"]["paragraphs"])
    letter_words, word_problem, sign_off = check_letter_words(spec["letter"]["paragraphs"], spec.get("letterWords"), spec.get("letterRevised") is True)
    if word_problem:
        problems.append(word_problem)

    build_cv(spec, cv)
    build_letter(spec, letter)
    cv_pdf, letter_pdf = to_pdf(cv), to_pdf(letter)

    cv_pages = pdf_pages(cv_pdf)
    if len(cv_pages) != 2:
        problems.append(f"CV is {len(cv_pages)} pages, not 2: shorten the Summary or a highlight.")
    elif page_two not in cv_pages[1].strip().splitlines()[0] if cv_pages[1].strip() else True:
        problems.append(f"Page two of the CV does not start with the {page_two} entry: page one has overflowed.")
    letter_pages = pdf_pages(letter_pdf)
    if len(letter_pages) != 1:
        problems.append(f"Letter is {len(letter_pages)} pages, not 1: cut it.")

    cv.with_suffix(".md").write_text(docx_to_md(cv), encoding="utf-8")
    letter.with_suffix(".md").write_text(docx_to_md(letter), encoding="utf-8")
    # The words as built, so tools/pack_diff.py can later show what the owner
    # changed before sending (backlog F2).
    baseline = out / "baseline.json"
    baseline.write_text(json.dumps({"cv": paragraphs(cv), "letter": paragraphs(letter)}, indent=1, ensure_ascii=False), encoding="utf-8")

    result = {
        "ok": not problems,
        "problems": problems,
        "files": [str(p) for p in [cv, cv_pdf, cv.with_suffix(".md"), letter, letter_pdf, letter.with_suffix(".md"), baseline]],
        "cvPages": len(cv_pages),
        "letterPages": len(letter_pages),
        "letterWords": letter_words,
        "signOff": [sign_off] if sign_off else [],
    }
    print(json.dumps(result, indent=2))
    sys.exit(0 if result["ok"] else 1)


if __name__ == "__main__":
    main()
